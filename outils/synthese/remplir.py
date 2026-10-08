#!/usr/bin/env python3
"""Remplit l'annexe 6-1 (tableau de synthèse, épreuve E4) à partir du contenu du portfolio.

Usage (depuis la racine du dépôt) :
    python3 outils/synthese/remplir.py [sortie.xlsx]

Lit les réalisations avec `node outils/synthese/export.js`, part du modèle officiel vierge
(modele-annexe-6-1-E4.xlsx) et écrit le fichier rempli (par défaut
outils/synthese/sortie/Tableau-de-synthese-E4.xlsx). Seules les compétences du bloc 1 figurent
dans ce tableau. Reste à compléter à la main : le n° de candidat et, si besoin, les dates
exactes au format JJ/MM/AA.
"""
import json, os, subprocess, sys
from copy import copy
import openpyxl
from openpyxl.styles import Alignment, Font

ICI = os.path.dirname(os.path.abspath(__file__))
DEPOT = os.path.dirname(os.path.dirname(ICI))
MODELE = os.path.join(ICI, 'modele-annexe-6-1-E4.xlsx')
SORTIE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ICI, 'sortie', 'Tableau-de-synthese-E4.xlsx')

# Lignes disponibles dans le modèle pour chaque rubrique.
LIGNES = {'formation': range(9, 19), 'pro1': range(20, 27), 'pro2': range(28, 35)}
COLONNES = ['C', 'D', 'E', 'F', 'G', 'H']  # même ordre que le bloc 1 de referentiel.js

data = json.loads(subprocess.check_output(['node', os.path.join(ICI, 'export.js')], cwd=DEPOT))
wb = openpyxl.load_workbook(MODELE)
ws = wb.active

ws['A8'] = 'Réalisations en cours de formation'  # le modèle a des retours à la ligne qui masquent ce titre
ws['A3'] = f"NOM et prénom : {data['nom']}"
ws['A4'] = f"Centre de formation : {data['centre']}"
ws['G4'] = '☒ SISR' if data['option'] == 'SISR' else '▢ SISR'
ws['H4'] = '☒ SLAM' if data['option'] == 'SLAM' else '▢ SLAM'
ws['A5'] = f"Adresse URL du portfolio : {data['url']}"

for rubrique, lignes in LIGNES.items():
    items = [r for r in data['realisations'] if r['rubrique'] == rubrique]
    if len(items) > len(lignes):
        sys.exit(f"Trop de réalisations pour la rubrique {rubrique} ({len(items)} pour {len(lignes)} lignes).")
    for r, ligne in zip(items, lignes):
        a = ws[f'A{ligne}']
        a.value = r['titre'] + ''.join(f"\n   – {d}" for d in r['documents'])
        a.alignment = Alignment(vertical='center', wrap_text=True)
        f = copy(a.font); f.sz = 10; a.font = f
        b = ws[f'B{ligne}']
        b.value = r['periode']
        b.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        fb = copy(b.font); fb.sz = 9; fb.b = False; b.font = fb
        for col, comp in zip(COLONNES, data['competences']):
            c = ws[f'{col}{ligne}']
            c.value = 'X' if comp in r['competences'] else None
            c.font = Font(name='Arial', sz=14, b=True)
            c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[ligne].height = max(40, 15 * (1 + len(r['documents'])) + 10)

os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
wb.save(SORTIE)
print(f"Tableau de synthèse écrit : {os.path.relpath(SORTIE, DEPOT)} ({len(data['realisations'])} réalisations)")
